"""Build small editable samples. All rows are synthetic and contain no personal data."""
import csv
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

DEST = Path(__file__).resolve().parents[1] / 'web' / 'public' / 'templates'
DEST.mkdir(parents=True, exist_ok=True)
questions = [
 ['meeting_01', 1, '请帮我列一个半小时项目例会的议程。', '提供可在半小时内完成的议程，包含问题和下一步。', '', '核心任务'],
 ['clarify_01', 1, '帮我把这件事安排一下。', '先询问要安排什么，不编造时间、地点或参与者。', '', '信息不足'],
 ['context_01', 1, '请帮我安排周三下午的项目讨论。', '保留周三下午，未确定事项应询问。', '', '多轮更正'],
 ['context_01', 2, '改成周四上午，帮我确认最新时间。', '采用周四上午，不再沿用周三下午。', '用户在本轮已明确更正为周四上午。', '多轮更正'],
]
answers = [[r[0],r[1],r[2],a,r[4]] for r,a in zip(questions,[
 '建议用五分钟确认目标、十五分钟讨论问题、十分钟明确下一步。',
 '你希望安排什么事情？可以先告诉我目标和时间要求。',
 '已记下周三下午。讨论需要多长时间？',
 '最新时间是周三下午。',
])]
for kind,headers,rows in [
 ('questions',['用例 ID','轮次','用户问题','希望做到什么','参考依据','场景'],questions),
 ('answers',['会话 ID','轮次','用户问题','Agent 回答','参考答案'],answers),
]:
 with (DEST/f'{kind}.csv').open('w',encoding='utf-8-sig',newline='') as f:
  writer=csv.writer(f);writer.writerow(headers);writer.writerows(rows)
 wb=Workbook();ws=wb.active;ws.title='填写材料';ws.append(headers)
 for row in rows:ws.append(row)
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='2466A8')
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
 for key in ['A','B']:ws.column_dimensions[key].width=20
 for key in ['C','D','E','F']:ws.column_dimensions[key].width=48
 ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
 ws['A1'].comment=Comment('本表含三个合成示例，含一组两轮对话。替换为自己的脱敏材料；问题模板预期须审核；回答模板第四行故意保留错误，便于理解。不要加入第二张数据表。','Agent EVO')
 wb.save(DEST/f'{kind}.xlsx')
execution=[{'case_id':r[0],'turn_index':r[1],'message':r[2],'synthetic':True} for r in questions]
(DEST/'execution-example.jsonl').write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in execution)+'\n',encoding='utf-8')
(DEST/'README.txt').write_text('questions：待提问的问题与待确认预期。answers：已取得的回答。execution-example：仅含可发给 Agent 的输入，未包含标准答案或隐藏断言。所有示例均为合成材料。\n',encoding='utf-8')
print('Created editable material templates:', DEST)
