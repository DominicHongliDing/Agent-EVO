from __future__ import annotations

import hashlib
import json
import mimetypes
import secrets
import sqlite3
import subprocess
import time
from pathlib import Path
from urllib.parse import urljoin, urlparse
from uuid import uuid4

import httpx
from docx import Document
from openpyxl import load_workbook

from agent_evo.config import Settings
from agent_evo.store import Store, dumps


TERMINAL = {"succeeded", "failed", "cancelled"}


class AgentRunError(Exception):
    def __init__(self, code: str, detail: str):
        super().__init__(detail)
        self.code = code


def source_fingerprint(root: Path, relative_dirs: tuple[str, ...]) -> dict:
    if not root.is_dir():
        return {"source_sha256": "unknown", "reason": "source_root_missing"}
    digest = hashlib.sha256()
    count = 0
    for relative in relative_dirs:
        base = root / relative
        if not base.exists():
            continue
        files = [base] if base.is_file() else list(base.rglob("*.py"))
        for path in sorted(files):
            if not path.is_file():
                continue
            digest.update(str(path.relative_to(root)).replace("\\", "/").encode())
            digest.update(hashlib.sha256(path.read_bytes()).digest())
            count += 1
    result = {"source_sha256": digest.hexdigest() if count else "unknown", "source_file_count": count}
    git_dir = root
    while git_dir != git_dir.parent and not (git_dir / ".git").exists():
        git_dir = git_dir.parent
    if (git_dir / ".git").exists():
        head = subprocess.run(["git", "rev-parse", "HEAD"], cwd=git_dir, capture_output=True, text=True, check=False)
        dirty = subprocess.run(["git", "status", "--porcelain"], cwd=git_dir, capture_output=True, text=True, check=False)
        result.update({"git_head": head.stdout.strip() if head.returncode == 0 else "unknown",
                       "git_dirty": bool(dirty.stdout.strip()) if dirty.returncode == 0 else None})
    return result


def _client(base_url: str, timeout: int) -> httpx.Client:
    parsed = urlparse(base_url)
    if parsed.scheme != "http" or parsed.hostname not in {"127.0.0.1", "localhost"}:
        raise AgentRunError("nonlocal_target", "Agent target must be local loopback")
    return httpx.Client(base_url=base_url, timeout=httpx.Timeout(float(timeout), connect=5.0),
                        follow_redirects=False, trust_env=False)


def _fixture_path(root: Path, relative: str) -> Path:
    path = (root / relative).resolve()
    if not path.is_file() or root.resolve() not in path.parents:
        raise AgentRunError("invalid_fixture", "Fixture must be a file under the fixture directory")
    return path


def _check_file(path: Path) -> dict:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
            sheets = []
            formula_count = 0
            for sheet in workbook:
                formula_count += sum(isinstance(cell.value, str) and cell.value.startswith("=")
                                     for row in sheet for cell in row)
                sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column})
            workbook.close()
            return {"opened": True, "kind": "xlsx", "sheets": sheets, "formula_count": formula_count}
        if suffix == ".docx":
            document = Document(path)
            return {"opened": True, "kind": "docx", "paragraph_count": len(document.paragraphs),
                    "table_count": len(document.tables)}
        return {"opened": True, "kind": suffix.lstrip(".") or "binary"}
    except Exception as exc:
        return {"opened": False, "kind": suffix.lstrip("."), "error_type": type(exc).__name__}


def _unihr_usage(db_path: Path | None, execution_id: str) -> dict:
    if db_path is None or not db_path.is_file():
        return {"verified": False, "reason": "usage_database_unconfigured"}
    try:
        db = sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True, timeout=5)
        db.row_factory = sqlite3.Row
        try:
            rows = db.execute(
                "SELECT provider,model,purpose,operation_kind,status,input_units,output_units,"
                "latency_ms,estimated_cost_cny,pricing_configured,created_at "
                "FROM model_usage_events WHERE run_id=? ORDER BY created_at", (execution_id,),
            ).fetchall()
        finally:
            db.close()
        events = [dict(row) for row in rows]
        return {"verified": True, "calls": len([row for row in events if row["status"] == "succeeded"]),
                "bailian_calls": len([row for row in events if row["status"] == "succeeded" and row["provider"] == "bailian"]),
                "events": events}
    except (sqlite3.Error, OSError) as exc:
        return {"verified": False, "reason": type(exc).__name__}


def _safe_download_url(base_url: str, url: str) -> str:
    absolute = urljoin(base_url + "/", url)
    base = urlparse(base_url)
    parsed = urlparse(absolute)
    if (parsed.scheme, parsed.hostname, parsed.port) != (base.scheme, base.hostname, base.port) or \
       not parsed.path.startswith("/api/v1/artifacts/") or not parsed.path.endswith("/download"):
        raise AgentRunError("unsafe_artifact_url", "Artifact URL outside allowlisted business endpoint")
    return absolute


def run_unihr(run: dict, case: dict, settings: Settings, store: Store) -> dict:
    if not settings.unihr_username or not settings.unihr_password:
        raise AgentRunError("unihr_credentials_missing", "UniHR test account is not configured")
    if settings.unihr_usage_db is None:
        raise AgentRunError("unihr_usage_unconfigured", "UniHR usage database is required to verify live model calls")
    version = source_fingerprint(settings.unihr_source, ("app",))
    snapshot_path = settings.unihr_usage_db.parent / "knowledge-snapshot.json"
    if snapshot_path.is_file():
        snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
        version["knowledge_snapshot"] = snapshot
    store.set_run_fields(run["run_id"], version_json=dumps({**run.get("version", {}), **version}))
    store.add_event(run["run_id"], "adapter", "unihr.authenticate", "started", {})
    with _client(settings.unihr_url, settings.timeout_seconds) as client:
        login = client.post("/api/v1/auth/login", json={"username": settings.unihr_username,
                                                        "password": settings.unihr_password})
        if login.status_code != 200:
            raise AgentRunError("unihr_auth_failed", f"UniHR login returned HTTP {login.status_code}")
        token = login.json().get("access_token")
        if not token:
            raise AgentRunError("unihr_auth_failed", "UniHR login response had no token")
        client.headers["Authorization"] = f"Bearer {token}"
        store.add_event(run["run_id"], "adapter", "unihr.authenticate", "completed", {"http_status": 200})

        files = []
        handles = []
        try:
            for relative in case.get("fixtures", []):
                path = _fixture_path(settings.fixture_root, relative)
                handle = path.open("rb")
                handles.append(handle)
                files.append(("files", (path.name, handle, mimetypes.guess_type(path.name)[0] or "application/octet-stream")))
            store.add_event(run["run_id"], "adapter", "unihr.submit", "started",
                            {"endpoint": "/api/v1/assistant/turns:start", "fixture_count": len(files)})
            form = {"message": case["message"], "task_mode": case.get("expected", {}).get("task_mode", "自动选择")}
            if case.get("_conversation_id"):
                form["conversation_id"] = case["_conversation_id"]
            response = client.post("/api/v1/assistant/turns:start",
                                   data=form,
                                   files=files)
        finally:
            for handle in handles:
                handle.close()
        if response.status_code != 202:
            raise AgentRunError("unihr_submit_failed", f"UniHR submit returned HTTP {response.status_code}")
        current = response.json()
        execution_id = str(current.get("id") or "")
        if not execution_id:
            raise AgentRunError("unihr_execution_missing", "UniHR response lacked execution ID")
        store.set_run_fields(run["run_id"], business_execution_id=execution_id)
        store.add_event(run["run_id"], "adapter", "unihr.submit", "completed", {"execution_id": execution_id})
        cursor = 0
        deadline = time.monotonic() + settings.timeout_seconds
        while time.monotonic() < deadline:
            with client.stream("GET", f"/api/v1/assistant/turns/{execution_id}/events",
                               params={"after": cursor}) as stream:
                if stream.status_code != 200:
                    raise AgentRunError("unihr_events_failed", f"UniHR events returned HTTP {stream.status_code}")
                for raw in stream.iter_lines():
                    if time.monotonic() >= deadline:
                        break
                    if not raw.startswith("data: "):
                        continue
                    try:
                        item = json.loads(raw[6:])
                    except json.JSONDecodeError:
                        continue
                    cursor = max(cursor, int(item.get("sequence") or 0))
                    store.add_event(run["run_id"], "business_event", str(item.get("event") or "unihr.event"),
                                    str(item.get("status") or "unknown"), item)
            check = client.get(f"/api/v1/assistant/turns/{execution_id}")
            if check.status_code != 200:
                raise AgentRunError("unihr_poll_failed", f"UniHR poll returned HTTP {check.status_code}")
            current = check.json()
            if current.get("run_id"):
                store.set_run_fields(run["run_id"], business_run_id=str(current["run_id"]))
            if str(current.get("status")) in TERMINAL:
                break
            time.sleep(1)
        else:
            store.add_event(run["run_id"], "adapter", "unihr.timeout", "timed_out", {"execution_id": execution_id})
            raise AgentRunError("unihr_timeout", "UniHR execution exceeded deadline; business state must be checked before rerun")

        if current.get("run_id"):
            detail = client.get(f"/api/v1/runs/{current['run_id']}")
            if detail.status_code == 200:
                for item in detail.json().get("events", []):
                    store.add_event(run["run_id"], "business_step", str(item.get("event") or item.get("type") or "run.event"),
                                    str(item.get("status") or "unknown"), item)
        result = current.get("result") or {}
        for item in result.get("trace") or []:
            store.add_event(run["run_id"], "business_trace", str(item.get("event") or "unihr.trace"),
                            str(item.get("status") or "unknown"), item)
        saved_artifacts = []
        for item in result.get("artifacts") or []:
            source_id = str(item.get("id") or "")
            url = str(item.get("download_url") or f"/api/v1/artifacts/{source_id}/download")
            target_url = _safe_download_url(settings.unihr_url, url)
            downloaded = client.get(target_url)
            if downloaded.status_code != 200:
                store.add_event(run["run_id"], "artifact", "unihr.download", "failed",
                                {"source_id": source_id, "http_status": downloaded.status_code})
                continue
            filename = Path(str(item.get("name") or source_id)).name
            target_dir = settings.artifact_root / run["run_id"]
            target_dir.mkdir(parents=True, exist_ok=True)
            target = target_dir / f"{hashlib.sha256(source_id.encode()).hexdigest()[:16]}-{filename}"
            target.write_bytes(downloaded.content)
            checks = _check_file(target)
            record = store.add_artifact(run["run_id"], source_id, filename,
                                        str(item.get("mime_type") or "application/octet-stream"), target, checks)
            saved_artifacts.append({k: v for k, v in record.items() if k != "local_path"})
            store.add_event(run["run_id"], "artifact", "unihr.download", "completed",
                            {"artifact_id": record["artifact_id"], "sha256": record["sha256"], "checks": checks})
        usage = _unihr_usage(settings.unihr_usage_db, execution_id)
        store.set_run_fields(run["run_id"], model_calls=int(usage.get("calls", 0)),
                             model_evidence_json=dumps(usage))
        return {"business_status": current.get("status"), "error_code": current.get("error_code"),
                "result": result, "artifacts": saved_artifacts, "model_evidence": usage,
                "conversation_id": current.get("conversation_id") or (result.get("conversation") or {}).get("id")}


def _prepare_clinical_identity(client: httpx.Client, run: dict,
                               settings: Settings, store: Store) -> str:
    token = settings.clinical_access_token
    if not token:
        username = f"evo_{uuid4().hex[:18]}"
        password = secrets.token_urlsafe(24)
        registered = client.post("/auth/register", json={"username": username, "password": password,
                                                        "preferred_language": "zh"})
        if registered.status_code not in {200, 201}:
            raise AgentRunError("clinical_context_auth_failed", "Synthetic clinician registration failed; configure EVO_CLINICAL_ACCESS_TOKEN")
        login = client.post("/auth/login", json={"username": username, "password": password})
        if login.status_code != 200:
            raise AgentRunError("clinical_context_auth_failed", "Synthetic clinician login failed")
        token = str(login.json().get("access_token") or "")
    client.headers["Authorization"] = f"Bearer {token}"
    me = client.get("/auth/me")
    if me.status_code != 200:
        raise AgentRunError("clinical_context_auth_failed", "Synthetic clinician token is invalid")
    user_id = str((me.json().get("user") or {}).get("user_id") or "")
    if not user_id:
        raise AgentRunError("clinical_context_auth_failed", "Synthetic clinician ID is missing")
    store.add_event(run["run_id"], "adapter", "clinical.authenticate", "completed", {"user_id": user_id})
    return user_id


def _prepare_clinical_context(client: httpx.Client, run: dict, case: dict,
                              settings: Settings, store: Store, thread_id: str) -> tuple[str, str, dict]:
    user_id = _prepare_clinical_identity(client, run, settings, store)
    created = client.post("/cases", json={"title": case["title"], "patient_alias": f"SYN-{run['run_id'][-8:]}"})
    if created.status_code not in {200, 201}:
        raise AgentRunError("clinical_case_create_failed", f"Clinical case creation returned HTTP {created.status_code}")
    case_id = str(created.json().get("case_id") or "")
    if not case_id:
        raise AgentRunError("clinical_case_create_failed", "Clinical case ID missing")
    context_text = str(case["case_context_text"])
    imported = client.post(f"/cases/{case_id}/import-text", json={"text": context_text})
    if imported.status_code != 200:
        raise AgentRunError("clinical_case_import_failed", f"Clinical case import returned HTTP {imported.status_code}")
    store.add_event(run["run_id"], "adapter", "clinical.case_context", "completed",
                    {"case_id": case_id, "context_sha256": hashlib.sha256(context_text.encode()).hexdigest()})
    setup_request = {"thread_id": thread_id, "case_id": case_id, "user_id": user_id, "language": "zh",
                     "message": "请整理这份病例记录中的已知事实、信息来源和待确认项，方便后续问诊；说明相关依据。"}
    setup = client.post("/chat", json=setup_request)
    if setup.status_code != 200:
        raise AgentRunError("clinical_setup_turn_failed", f"Clinical setup turn returned HTTP {setup.status_code}")
    setup_result = setup.json()
    setup_usage = (setup_result.get("trace") or {}).get("model_usage") or {}
    setup_calls = _clinical_call_count(setup_usage)
    store.add_event(run["run_id"], "business_setup", "clinical.setup_turn", "completed", setup_result)
    store.set_run_fields(run["run_id"], model_calls=setup_calls,
                         model_evidence_json=dumps({"verified": False, "setup_calls": setup_calls,
                                                    "setup_model_usage": setup_usage}))
    return user_id, case_id, {"run_id": setup_result.get("run_id"), "model_usage": setup_usage,
                              "calls": setup_calls}


def run_clinical(run: dict, case: dict, settings: Settings, store: Store, session: dict | None = None) -> dict:
    version = source_fingerprint(settings.clinical_source, ("src/clinical_rag",))
    corpus = settings.clinical_source / "data" / "processed" / "chunks" / "clinical_rag_chunks.jsonl"
    version["local_corpus_sha256"] = hashlib.sha256(corpus.read_bytes()).hexdigest() if corpus.is_file() else "unknown"
    store.set_run_fields(run["run_id"], version_json=dumps({**run.get("version", {}), **version}))
    # One unique synthetic thread per platform run permits evidence lookup after a broken stream.
    session = session if session is not None else {}
    thread_id = session.setdefault("thread_id", f"evo_{run['run_id']}")
    final = None
    stream_error: Exception | None = None
    with _client(settings.clinical_url, settings.timeout_seconds) as client:
        case_id = None
        setup = None
        if session.get("user_id"):
            user_id, case_id = session["user_id"], session.get("case_id")
            client.headers["Authorization"] = session["token"]
        elif case.get("case_context_text"):
            user_id, case_id, setup = _prepare_clinical_context(client, run, case, settings, store, thread_id)
        else:
            user_id = _prepare_clinical_identity(client, run, settings, store)
        session.update(user_id=user_id, case_id=case_id, token=client.headers["Authorization"])
        request = {"thread_id": thread_id, "user_id": user_id,
                   "message": case["message"], "language": "zh"}
        if case_id:
            request["case_id"] = case_id
        store.add_event(run["run_id"], "adapter", "clinical.submit", "started",
                        {"endpoint": "/chat/stream", "thread_id": thread_id, "case_id": case_id})
        try:
            with client.stream("POST", "/chat/stream", json=request) as response:
                if response.status_code != 200:
                    raise AgentRunError("clinical_submit_failed", f"Clinical returned HTTP {response.status_code}")
                for line in response.iter_lines():
                    if not line.startswith("data: "):
                        continue
                    try:
                        event = json.loads(line[6:])
                    except json.JSONDecodeError:
                        continue
                    event_type = str(event.get("type") or "clinical.event")
                    store.add_event(run["run_id"], "business_event", event_type, "observed", event)
                    if event_type == "final":
                        final = event.get("response") or {}
        except (httpx.HTTPError, AgentRunError) as exc:
            stream_error = exc
            store.add_event(run["run_id"], "adapter", "clinical.stream", "failed",
                            {"error_type": type(exc).__name__, "error_code": getattr(exc, "code", None)})
        business_id = str(final.get("run_id") or "") if isinstance(final, dict) else ""
        business_row = None
        listed = client.get("/admin/agent-runs")
        if listed.status_code == 200:
            matches = [item for item in listed.json().get("items", [])
                       if item.get("thread_id") == thread_id and item.get("user_id") == user_id
                       and item.get("run_id") != (setup or {}).get("run_id")
                       and item.get("run_id") not in session.get("previous_business_ids", [])]
            if business_id:
                matches = [item for item in matches if item.get("run_id") == business_id]
            if matches:
                business_row = matches[0]
                business_id = business_id or str(business_row.get("run_id") or "")
                snapshot = business_row.get("version_snapshot") or {}
                if isinstance(snapshot, dict):
                    store.set_run_fields(run["run_id"], version_json=dumps({**run.get("version", {}), **version,
                                                                             "business_version_snapshot": snapshot}))
        if business_id:
            store.set_run_fields(run["run_id"], business_run_id=business_id)
            from agent_evo.evidence import import_model_evidence
            import_model_evidence(store, run['run_id'], settings.clinical_source, business_id)
            steps = client.get(f"/admin/agent-runs/{business_id}/steps")
            if steps.status_code == 200:
                for step in steps.json().get("items", []):
                    store.add_event(run["run_id"], "business_step", str(step.get("node_name") or "clinical.step"),
                                    str(step.get("status") or "unknown"), step)
        trace = (final.get("trace") or {}) if isinstance(final, dict) else {}
        usage = trace.get("model_usage") or ((business_row or {}).get("final_trace") or {}).get("model_usage") or {}
        main_calls = _clinical_call_count(usage)
        setup_calls = int((setup or {}).get("calls") or 0)
        evidence = {"verified": bool(business_id) and main_calls > 0, "calls": main_calls + setup_calls,
                    "main_calls": main_calls, "setup_calls": setup_calls,
                    "model_usage": usage, "setup_model_usage": (setup or {}).get("model_usage"),
                    "business_run_id": business_id or None, "setup_business_run_id": (setup or {}).get("run_id")}
        store.set_run_fields(run["run_id"], model_calls=main_calls + setup_calls, model_evidence_json=dumps(evidence))
        if final is None:
            if isinstance(stream_error, httpx.TimeoutException):
                raise AgentRunError("clinical_timeout", "Clinical stream timed out; inspect persisted run before rerun")
            raise AgentRunError("clinical_no_final", "Clinical stream ended without final response; inspect persisted steps")
        session.setdefault("previous_business_ids", []).append(business_id)
        return {"business_status": "succeeded", "result": final, "model_evidence": evidence,
                "conversation_id": final.get("thread_id")}


def _clinical_call_count(usage: object) -> int:
    if not isinstance(usage, dict):
        return 0
    for key in ("successful_request_count", "call_count", "calls", "total_calls"):
        value = usage.get(key)
        if isinstance(value, (int, float)):
            return int(value)
    models = usage.get("models")
    if isinstance(models, list):
        return len(models)
    return 0
