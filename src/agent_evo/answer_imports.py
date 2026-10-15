"""Immutable answer files, explicit mapping, and independent answer evaluation.

This module never calls an Agent or creates a business run. Imported evidence is
user-provided; it is not proof that a model or a tool actually ran.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import sqlite3
import zipfile
from collections import OrderedDict
from pathlib import Path
from uuid import uuid4

from agent_evo.store import dumps, now

LIMIT = 5 * 1024 * 1024
FIELDS = {
    'question': ['用户问题', '问题', '用户输入', 'question', 'input', 'prompt'],
    'answer': ['agent回答', 'agent 回答', '回答', '回复', 'answer', 'output', 'response', 'prediction', 'pred'],
    'reference': ['参考答案', '期望回答', '期望', 'reference', 'expected_output', 'gold_answer'],
    'conversation': ['会话id', '会话 id', 'conversation_id', 'session_id', 'conversation'],
    'turn': ['轮次', 'turn', 'turn_index', 'round'],
    'case_id': ['用例id', '用例 id', 'case_id', 'question_id'],
    'version': ['测试集版本', 'dataset_version', 'version'],
    'trace': ['执行轨迹', 'trace', 'events'],
}


def digest(value):
    return hashlib.sha256(dumps(value).encode()).hexdigest()


def migrate(store):
    with store.connect() as db:
        db.executescript('''
        CREATE TABLE IF NOT EXISTS answer_imports (
          import_id TEXT PRIMARY KEY, filename TEXT NOT NULL, sha256 TEXT NOT NULL,
          raw_json TEXT NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS answer_batches (
          batch_id TEXT PRIMARY KEY, import_id TEXT NOT NULL REFERENCES answer_imports(import_id),
          fingerprint TEXT NOT NULL UNIQUE, name TEXT NOT NULL, mapping_json TEXT NOT NULL,
          created_at TEXT NOT NULL, deleted_at TEXT, pinned INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS answer_records (
          record_id TEXT PRIMARY KEY, batch_id TEXT NOT NULL REFERENCES answer_batches(batch_id),
          ordinal INTEGER NOT NULL, data_json TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS answer_assessments (
          assessment_id TEXT PRIMARY KEY, record_id TEXT NOT NULL REFERENCES answer_records(record_id),
          status TEXT NOT NULL, created_at TEXT NOT NULL, finished_at TEXT, result_json TEXT NOT NULL DEFAULT '{}'
        );
        CREATE INDEX IF NOT EXISTS answer_jobs ON answer_assessments(status, created_at);
        CREATE INDEX IF NOT EXISTS answer_history ON answer_assessments(record_id, created_at);
        ''')


def clean_text(value):
    if value is None:
        return ''
    if not isinstance(value, (str, int, float)) or isinstance(value, bool):
        raise ValueError('问答、轮次和标识字段须为文本或数字；多轮请使用 turns / messages 数组')
    return str(value).strip()


def private_material(value):
    text = dumps(value)
    return bool(re.search(r'"(?:password|access_token|api_key|secret|authorization)"\s*:|\bsk-[A-Za-z0-9_-]{16,}|(?<!\d)1[3-9]\d{9}(?!\d)|(?<!\d)\d{17}[\dXx](?!\d)|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', text, re.I))


def read_json(text):
    try:
        value = json.loads(text)
    except RecursionError:
        raise ValueError('JSON 嵌套层级过多，请使用问答记录数组') from None
    def check(item):
        if isinstance(item, float) and not math.isfinite(item):
            raise ValueError('JSON 含非有限数值，请改成可读文本或有效数字')
        if isinstance(item, dict):
            for v in item.values():
                check(v)
        if isinstance(item, list):
            for v in item:
                check(v)
    try:
        check(value)
    except RecursionError:
        raise ValueError('JSON 嵌套层级过多，请使用问答记录数组') from None
    return value


def parse_file(filename: str, content: bytes):
    if not content or len(content) > LIMIT:
        raise ValueError('请选择非空文件，大小不超过 5 MiB')
    suffix = Path(filename).suffix.lower()
    errors, rows = [], []
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('Excel 解压后过大，请拆分文件')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            filled = [sheet for sheet in book if sheet.max_row and sheet.max_row > 1]
            if len(filled) != 1:
                raise ValueError('请仅保留一个含数据的工作表后上传，避免漏读其他表')
            sheet = filled[0]
            values = sheet.iter_rows()
            headers = [clean_text(c.value) for c in next(values)]
            if not all(headers) or len(set(headers)) != len(headers):
                raise ValueError('表头含空列名或重复列名，请修正后上传')
            for line, cells in enumerate(values, 2):
                if line > 1001:
                    raise ValueError('一次最多导入 1,000 行，请拆分文件')
                if all(c.value is None for c in cells):
                    continue
                if any(c.data_type == 'f' for c in cells):
                    errors.append({'line': line, 'message': '含公式，请粘贴为值后导入；平台不执行公式'})
                    continue
                rows.append({'line': line, 'value': dict(zip(headers, [c.value for c in cells]))})
            book.close()
        except (zipfile.BadZipFile, KeyError, StopIteration) as exc:
            raise ValueError('Excel 文件不完整或没有有效表头') from exc
    else:
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = content.decode('gb18030') if suffix == '.csv' else content.decode('utf-8')
            except UnicodeDecodeError as exc:
                raise ValueError('请使用 UTF-8 文本，CSV 也支持 GB18030 编码') from exc
        if suffix == '.csv':
            try:
                reader = csv.DictReader(io.StringIO(text), dialect=csv.Sniffer().sniff(text[:8192], delimiters=',;\t'))
            except csv.Error:
                reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise ValueError('CSV 需要唯一的表头')
            try:
                for row in reader:
                    if None in row or None in row.values():
                        errors.append({'line': reader.line_num, 'message': '列数与表头不一致'})
                    else:
                        rows.append({'line': reader.line_num, 'value': row})
            except csv.Error as exc:
                raise ValueError('CSV 引号或分隔符格式不完整') from exc
        elif suffix in {'.json', '.jsonl'}:
            if suffix == '.jsonl':
                for n, line in enumerate(text.splitlines(), 1):
                    if line.strip():
                        try:
                            rows.append({'line': n, 'value': read_json(line)})
                        except ValueError:
                            errors.append({'line': n, 'message': '本行 JSON 格式不完整，请检查引号、逗号和括号'})
            else:
                try:
                    value = read_json(text)
                except json.JSONDecodeError as exc:
                    raise ValueError(f'JSON 第 {exc.lineno} 行、第 {exc.colno} 列格式不完整') from exc
                except RecursionError:
                    raise ValueError('JSON 嵌套层级过多，请使用问答记录数组') from None
                if private_material(value):
                    raise ValueError('文件发现可能的个人标识或凭证，请先脱敏再导入')
                if isinstance(value, dict):
                    value = value.get('records', value.get('items', [value]))
                if not isinstance(value, list):
                    raise ValueError('JSON 须为问答数组，或包含 records / items 数组')
                rows = [{'line': n, 'value': row} for n, row in enumerate(value, 1)]
        elif suffix == '.txt':
            matches = list(re.finditer(r'(?m)^(用户|User|助手|Assistant|Agent)[：:]\s*', text, re.I))
            messages = [{'role': 'user' if m[1].lower() in {'用户', 'user'} else 'assistant',
                         'content': text[m.end():matches[i+1].start() if i+1 < len(matches) else len(text)].strip()}
                        for i, m in enumerate(matches)]
            if not matches or text[:matches[0].start()].strip():
                raise ValueError('粘贴对话请以“用户：”和“助手：”分行标记，可连续粘贴多轮')
            rows = [{'line': 1, 'value': {'messages': messages}}]
        else:
            raise ValueError('支持 Excel (.xlsx)、CSV、JSON、JSONL 或标记角色的文本对话')
    if not 1 <= len(rows) <= 1000 and not errors:
        raise ValueError('一次须包含 1–1,000 行问答')
    if len(rows) > 1000:
        raise ValueError('一次最多导入 1,000 行，请拆分文件')
    expanded = []
    for row in rows:
        value, line = row['value'], row['line']
        if not isinstance(value, dict):
            errors.append({'line': line, 'message': '每条记录须为包含问题和回答的对象'})
            continue
        if private_material(value):
            raise ValueError(f'第 {line} 行发现可能的个人标识或凭证，请先脱敏再导入')
        if 'messages' in value or 'turns' in value:
            turns = value.get('turns')
            if 'messages' in value:
                messages = value['messages']
                if not isinstance(messages, list) or not messages or len(messages) % 2 or any(
                    not isinstance(m, dict) or m.get('role') != ('user' if i % 2 == 0 else 'assistant') or not isinstance(m.get('content'), str)
                    for i, m in enumerate(messages)):
                    errors.append({'line': line, 'message': 'messages 须按 user / assistant 成对排列；工具记录请放在 trace 字段'})
                    continue
                turns = [{'question': messages[i]['content'], 'answer': messages[i+1]['content']} for i in range(0, len(messages), 2)]
            if not isinstance(turns, list) or not turns or any(not isinstance(t, dict) for t in turns):
                errors.append({'line': line, 'message': 'turns 须为包含 question / answer 的数组'})
                continue
            for i, turn in enumerate(turns, 1):
                expanded.append({'line': line, 'value': {**{k: v for k, v in value.items() if k not in {'turns', 'messages'}},
                    **turn, 'conversation_id': str(value.get('conversation_id') or f'file-row-{line}'), 'turn': i}})
        else:
            expanded.append(row)
    if len(expanded) > 1000:
        raise ValueError('展开多轮后超过 1,000 行，请拆分文件')
    expanded_size = 0
    for row in expanded:
        expanded_size += len(dumps(row).encode())
        if expanded_size > 2 * LIMIT:
            raise ValueError('展开多轮后的材料超过 10 MiB，请减少重复轨迹或拆分文件')
    columns = list(dict.fromkeys(k for row in expanded for k in row['value']))
    mapping, ambiguous = {}, []
    for field, aliases in FIELDS.items():
        options = [c for c in columns if c.strip().lower() in aliases]
        mapping[field] = options[0] if len(options) == 1 else ''
        if len(options) > 1:
            ambiguous.append(field)
    return {'rows': expanded, 'columns': columns, 'mapping': mapping, 'ambiguous': ambiguous, 'errors': errors}


def create_import(store, filename, content):
    raw = parse_file(filename, content)
    identifier = 'import_' + uuid4().hex
    folder = store.settings.artifact_root / 'answer-imports'
    folder.mkdir(parents=True, exist_ok=True)
    (folder / (identifier + Path(filename).suffix.lower())).write_bytes(content)
    with store.connect() as db:
        db.execute('INSERT INTO answer_imports VALUES(?,?,?,?,?)',
                   (identifier, Path(filename).name[:200], hashlib.sha256(content).hexdigest(), dumps(raw), now()))
    return import_info(store, identifier)


def import_info(store, identifier):
    with store.connect() as db:
        row = db.execute('SELECT * FROM answer_imports WHERE import_id=?', (identifier,)).fetchone()
    if not row:
        raise KeyError('导入文件不存在')
    raw = json.loads(row['raw_json'])
    return {'import_id': identifier, 'filename': row['filename'], 'sha256': row['sha256'], **raw}


def normalize(store, identifier, mapping, dataset_id='', confirm_order=False):
    raw = import_info(store, identifier)
    if set(mapping) - set(FIELDS) or any(c and c not in raw['columns'] for c in mapping.values()):
        raise ValueError('列映射无效，请重新选择')
    used = [c for c in mapping.values() if c]
    errors = list(raw['errors'])
    if len(set(used)) != len(used):
        errors.append({'line': 0, 'message': '同一列不能同时作为不同字段'})
    if not mapping.get('answer'):
        errors.append({'line': 0, 'message': '请选择 Agent 回答所在列'})
    if not mapping.get('question') and not (dataset_id and mapping.get('case_id')):
        errors.append({'line': 0, 'message': '请选择用户问题列，或通过用例 ID 关联测试集'})
    dataset, lookup = None, {}
    if dataset_id:
        dataset = store.get_dataset(dataset_id)
        if dataset['profile'] not in {'clinical', 'generic'} or dataset['deleted_at']:
            raise ValueError('所选测试集不可用')
        for case in store.list_cases(dataset['profile'], dataset_id):
            for key in {case['case_id'], str(case.get('source_case_id', ''))} - {''}:
                lookup.setdefault(key, []).append(case)
    grouped = OrderedDict()
    matches = {'matched': 0, 'unmatched': 0, 'conflict': 0}
    for row in raw['rows']:
        line, value = row['line'], row['value']
        try:
            fields = {f: clean_text(value.get(mapping.get(f, ''))) for f in FIELDS if f != 'trace'}
            if not fields['answer']:
                raise ValueError('Agent 回答为空，请补齐或移除该行后重新上传')
            key = fields['conversation'] or f'line-{line}'
            group = grouped.setdefault(key, {'conversation_id': fields['conversation'], 'turns': [], 'source_lines': [], 'case_id': fields['case_id']})
            turn = len(group['turns']) + 1
            if fields['turn']:
                if not re.fullmatch(r'[1-9]\d*', fields['turn']):
                    raise ValueError('轮次须为从 1 开始的整数')
                turn = int(fields['turn'])
            elif fields['conversation'] and not confirm_order:
                raise ValueError('有会话 ID 但缺少轮次，请确认按文件顺序排列')
            if any(t['turn'] == turn for t in group['turns']):
                raise ValueError('同一会话出现重复轮次')
            if fields['case_id'] != group['case_id']:
                raise ValueError('同一会话的用例 ID 不一致')
            candidates = lookup.get(fields['case_id'], [])
            linked = None
            if dataset:
                if len(candidates) > 1:
                    raise ValueError('用例 ID 在所选测试集中不唯一，请使用平台 case_id')
                case = candidates[0] if candidates else None
                if case:
                    expected_questions = [t['message'] for t in case.get('turns', [])] or [case['message']]
                    if fields['version'] and fields['version'] != dataset['version']:
                        matches['conflict'] += 1
                        raise ValueError('测试集版本不一致，不能自动关联')
                    if turn > len(expected_questions) or (fields['question'] and fields['question'] != expected_questions[turn-1]):
                        matches['conflict'] += 1
                        raise ValueError('问题内容与测试集不一致，不能按用例 ID 覆盖关联')
                    if not fields['question'] and fields['version'] != dataset['version']:
                        raise ValueError('仅凭 ID 补问题时还需提供匹配的测试集版本')
                    fields['question'] = expected_questions[turn-1]
                    expected = case.get('expected', {})
                    linked = {'case_id': case['case_id'], 'dataset_id': dataset_id, 'version': dataset['version'],
                              'sha256': dataset['content_sha256'], 'expected': expected, 'gold_status': case['gold_status'],
                              'turn_checks': case.get('turns', [])[turn-1].get('checks', []) if case.get('turns') else []}
                    matches['matched'] += 1
                else:
                    matches['unmatched'] += 1
            if not fields['question']:
                raise ValueError('用户问题为空，且未能可靠关联原用例')
            if any(len(fields[k]) > 40000 for k in ('question', 'answer', 'reference')):
                raise ValueError('单个问答字段超过 40,000 字符，请拆分或人工复核；原文未截断')
            trace = value.get(mapping.get('trace', ''), [])
            if isinstance(trace, str) and trace.strip():
                try:
                    trace = read_json(trace)
                except ValueError:
                    trace = {'text': trace}
            group['turns'].append({'turn': turn, 'question': fields['question'], 'answer': fields['answer'],
                'reference': fields['reference'], 'line': line, 'linked_case': linked, 'trace': trace or []})
            group['source_lines'].append(line)
        except (ValueError, TypeError) as exc:
            errors.append({'line': line, 'message': str(exc)})
    records, seen_cases = [], set()
    for group in grouped.values():
        if not group['turns']:
            continue
        group['turns'].sort(key=lambda t: t['turn'])
        if [t['turn'] for t in group['turns']] != list(range(1, len(group['turns'])+1)):
            errors.append({'line': group['source_lines'][0], 'message': '会话轮次不连续或缺少首轮，请补全后导入'})
        if group['case_id'] and not group['conversation_id'] and group['case_id'] in seen_cases:
            errors.append({'line': group['source_lines'][0], 'message': '重复用例 ID；多轮或多次执行请提供不同的会话 ID'})
        seen_cases.add(group['case_id'])
        if len(group['turns']) > 40:
            errors.append({'line': group['source_lines'][0], 'message': '单段对话最多 40 轮，请拆分并保留必要上下文'})
        group['source'] = {'type': 'uploaded_answers', 'import_id': identifier, 'filename': raw['filename'], 'sha256': raw['sha256'],
                           'business_calls': 0, 'execution_verified': False}
        records.append(group)
    if not records:
        errors.append({'line': 0, 'message': '没有可评测的完整问答'})
    if len(records) > 100:
        errors.append({'line': 0, 'message': '单批最多 100 段对话，请拆分文件'})
    from agent_evo.answer_judge import segments
    return {'valid': not errors, 'errors': errors, 'records': records, 'matches': matches,
            'judge_segments': sum(len(segments(r)) for r in records),
            'total': len(records), 'turn_count': sum(len(r['turns']) for r in records)}


def commit_import(store, identifier, payload):
    if not payload.reviewed_deidentified:
        raise ValueError('请确认内容为合成或已脱敏的问答')
    parsed = normalize(store, identifier, payload.mapping, payload.dataset_id, payload.confirm_order)
    if not parsed['valid']:
        raise ValueError('仍有逐行问题，请返回预览修正')
    # A new upload id is not a new set of answers. Duplicate submissions retain
    # the original batch and do not queue another billed assessment.
    fingerprint = digest([{**r, 'source': {k: v for k, v in r['source'].items() if k != 'import_id'}} for r in parsed['records']])
    with store.connect() as db:
        db.execute('BEGIN IMMEDIATE')
        existing = db.execute('SELECT batch_id,deleted_at FROM answer_batches WHERE fingerprint=?', (fingerprint,)).fetchone()
        if existing:
            if existing['deleted_at']:
                raise ValueError('相同内容已在回收站，请先恢复原批次')
            return {'batch_id': existing['batch_id'], 'reused': True}
        batch_id = 'answers_' + uuid4().hex
        db.execute('INSERT INTO answer_batches(batch_id,import_id,fingerprint,name,mapping_json,created_at) VALUES(?,?,?,?,?,?)',
                   (batch_id, identifier, fingerprint, payload.name.strip() or '导入回答', dumps(payload.mapping), now()))
        for i, record in enumerate(parsed['records']):
            db.execute('INSERT INTO answer_records VALUES(?,?,?,?)', ('answer_' + uuid4().hex, batch_id, i, dumps(record)))
    return {'batch_id': batch_id, 'reused': False}


def record_detail(store, identifier):
    with store.connect() as db:
        row = db.execute('SELECT * FROM answer_records WHERE record_id=?', (identifier,)).fetchone()
        if not row:
            raise KeyError('回答记录不存在')
        history = [dict(r) for r in db.execute('SELECT * FROM answer_assessments WHERE record_id=? ORDER BY rowid DESC', (identifier,))]
    data = json.loads(row['data_json'])
    source = data['source'].get('import_id')
    original_rows = [r for r in import_info(store, source)['rows'] if r['line'] in data['source_lines']] if source else []
    return {'record_id': identifier, 'batch_id': row['batch_id'], **data, 'original_fields': original_rows,
            'assessments': [{k: v for k, v in r.items() if k != 'result_json'} | {'result': json.loads(r['result_json'])} for r in history]}


def list_batches(store, deleted=False):
    with store.connect() as db:
        rows = [dict(r) for r in db.execute('SELECT * FROM answer_batches WHERE (deleted_at IS NOT NULL)=? ORDER BY pinned DESC,created_at DESC', (int(deleted),))]
        for row in rows:
            records = [dict(r) for r in db.execute('SELECT record_id,data_json FROM answer_records WHERE batch_id=? ORDER BY ordinal', (row['batch_id'],))]
            summaries = []
            for record in records:
                data = json.loads(record['data_json'])
                job = db.execute('SELECT status,result_json FROM answer_assessments WHERE record_id=? ORDER BY rowid DESC LIMIT 1', (record['record_id'],)).fetchone()
                result = json.loads(job['result_json']) if job else {}
                summaries.append({'record_id': record['record_id'], 'question': data['turns'][0]['question'],
                    'turn_count': len(data['turns']), 'status': job['status'] if job else 'not_started',
                    'quality': result.get('quality', 'unjudged'), 'critical_count': len(result.get('critical_findings', []))})
            row.pop('mapping_json')
            row['records'] = summaries
            row['total'] = len(summaries)
    return rows


def queue_assessments(store, batch_id, selected=None):
    with store.connect() as db:
        db.execute('BEGIN IMMEDIATE')
        batch = db.execute('SELECT * FROM answer_batches WHERE batch_id=? AND deleted_at IS NULL', (batch_id,)).fetchone()
        if not batch:
            raise KeyError('批次不存在或已删除')
        ids = [r[0] for r in db.execute('SELECT record_id FROM answer_records WHERE batch_id=?', (batch_id,))]
        if selected is not None:
            if not selected or len(set(selected)) != len(selected) or not set(selected) <= set(ids):
                raise ValueError('请选择本批次中的有效记录')
            ids = selected
        added = []
        for rid in ids:
            if db.execute("SELECT 1 FROM answer_assessments WHERE record_id=? AND status IN ('queued','running')", (rid,)).fetchone():
                continue
            aid = 'answer_judge_' + uuid4().hex
            db.execute('INSERT INTO answer_assessments(assessment_id,record_id,status,created_at) VALUES(?,?,?,?)', (aid, rid, 'queued', now()))
            added.append(aid)
    return {'queued': len(added), 'assessment_ids': added, 'business_calls': 0}
